from __future__ import annotations

import hashlib
import json
import logging
import os
import re
import shutil
import threading
import time
import traceback
from dataclasses import asdict, dataclass
from pathlib import Path, PurePosixPath
from typing import Any, Callable, Iterable, Protocol


APP_VERSION = "1.0.0"
ProgressCallback = Callable[[str, int, int], None]


@dataclass(frozen=True)
class JobConfig:
    input_pdf: Path
    output_dir: Path
    language: str = "en"
    device: str = "cpu"
    profile: str = "quality"
    make_markdown: bool = True
    make_docx: bool = True
    offline: bool = False
    retries: int = 1
    orientation: bool = True
    textline_orientation: bool = True
    unwarp: bool = False
    tables: bool = True
    formulas: bool = False
    charts: bool = False

    def normalized(self) -> "JobConfig":
        return JobConfig(
            input_pdf=self.input_pdf.expanduser().resolve(),
            output_dir=self.output_dir.expanduser().resolve(),
            language=(self.language or "en").strip(),
            device=(self.device or "cpu").strip(),
            profile=(self.profile or "quality").strip(),
            make_markdown=bool(self.make_markdown),
            make_docx=bool(self.make_docx),
            offline=bool(self.offline),
            retries=max(0, min(int(self.retries), 5)),
            orientation=bool(self.orientation),
            textline_orientation=bool(self.textline_orientation),
            unwarp=bool(self.unwarp),
            tables=bool(self.tables),
            formulas=bool(self.formulas),
            charts=bool(self.charts),
        )


@dataclass
class JobSummary:
    total_pages: int
    completed_pages: int
    failed_pages: list[int]
    skipped_pages: int
    markdown_path: Path | None
    docx_path: Path | None
    log_path: Path
    manifest_path: Path
    cancelled: bool = False


class PageResult(Protocol):
    @property
    def markdown(self) -> dict[str, Any]: ...

    def save_to_word(self, save_path: str) -> None: ...

    def save_to_json(self, save_path: str) -> None: ...


class OcrEngine(Protocol):
    def process_page(self, page_pdf: Path) -> PageResult: ...


class PaddleStructureEngine:
    """Thin, local-only adapter around PaddleOCR PP-StructureV3."""

    def __init__(self, config: JobConfig):
        if config.offline:
            os.environ["PADDLE_PDX_DISABLE_MODEL_SOURCE_CHECK"] = "True"
            os.environ["HF_HUB_OFFLINE"] = "1"
            os.environ["TRANSFORMERS_OFFLINE"] = "1"
        else:
            os.environ.setdefault("PADDLE_PDX_MODEL_SOURCE", "BOS")

        try:
            from paddleocr import PPStructureV3
        except ImportError as exc:
            raise RuntimeError(
                "PaddleOCR is not installed. Run Install Scan2Doc.bat first."
            ) from exc

        model_overrides: dict[str, str] = {}
        if config.profile == "fast_cpu":
            model_overrides["layout_detection_model_name"] = "PP-DocLayout-S"
            if config.language in {"en", "ch", "ja"}:
                model_overrides.update(
                    text_detection_model_name="PP-OCRv5_mobile_det",
                    text_recognition_model_name="PP-OCRv5_mobile_rec",
                )

        self.pipeline = PPStructureV3(
            lang=config.language,
            device=config.device,
            use_doc_orientation_classify=config.orientation,
            use_doc_unwarping=config.unwarp,
            use_textline_orientation=config.textline_orientation,
            use_table_recognition=config.tables,
            use_formula_recognition=config.formulas,
            use_chart_recognition=config.charts,
            use_seal_recognition=False,
            # Paddle 3.3.1's Windows oneDNN path can fail on PP-DocLayout
            # with ConvertPirAttribute2RuntimeAttribute. Keep it off by
            # default for dependable unattended runs; GPU is unaffected.
            enable_mkldnn=False,
            cpu_threads=max(1, min(os.cpu_count() or 4, 8)),
            **model_overrides,
        )

    def process_page(self, page_pdf: Path) -> PageResult:
        results = list(self.pipeline.predict_iter(input=str(page_pdf)))
        if len(results) != 1:
            raise RuntimeError(
                f"Expected one result for a one-page PDF, received {len(results)}."
            )
        return results[0]


def validate_input_pdf(path: Path) -> None:
    if not path.is_file():
        raise ValueError(f"PDF does not exist: {path}")
    if path.suffix.lower() != ".pdf":
        raise ValueError("Only local .pdf files are accepted.")
    with path.open("rb") as stream:
        if stream.read(5) != b"%PDF-":
            raise ValueError("The selected file does not have a valid PDF header.")


def file_sha256(path: Path, block_size: int = 4 * 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while block := stream.read(block_size):
            digest.update(block)
    return digest.hexdigest()


def config_fingerprint(config: JobConfig, input_digest: str) -> str:
    material = {
        "app_version": APP_VERSION,
        "input_sha256": input_digest,
        "language": config.language,
        "device": config.device,
        "profile": config.profile,
        "orientation": config.orientation,
        "textline_orientation": config.textline_orientation,
        "unwarp": config.unwarp,
        "tables": config.tables,
        "formulas": config.formulas,
        "charts": config.charts,
    }
    raw = json.dumps(material, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _atomic_json(path: Path, value: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + ".tmp")
    temp.write_text(json.dumps(value, indent=2, ensure_ascii=False), encoding="utf-8")
    os.replace(temp, path)


def _safe_name(value: str, fallback: str = "asset") -> str:
    name = re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("._")
    return name[:120] or fallback


def _write_markdown_assets(
    md_info: dict[str, Any], page_number: int, output_dir: Path
) -> str:
    text = str(md_info.get("markdown_texts") or md_info.get("text") or "")
    images = md_info.get("markdown_images") or md_info.get("images") or {}
    if not isinstance(images, dict):
        return text

    asset_dir = output_dir / "assets" / f"page_{page_number:04d}"
    for original, image in images.items():
        original_text = str(original).replace("\\", "/")
        base = _safe_name(PurePosixPath(original_text).name, "image.png")
        if "." not in base:
            base += ".png"
        asset_dir.mkdir(parents=True, exist_ok=True)
        destination = asset_dir / base
        suffix = 2
        while destination.exists():
            destination = asset_dir / f"{Path(base).stem}_{suffix}{Path(base).suffix}"
            suffix += 1
        if hasattr(image, "save"):
            image.save(destination)
        elif isinstance(image, (bytes, bytearray)):
            destination.write_bytes(bytes(image))
        else:
            logging.getLogger("scan2doc").warning(
                "Page %d: unsupported embedded image type for %s",
                page_number,
                original_text,
            )
            continue
        replacement = destination.relative_to(output_dir).as_posix()
        text = text.replace(original_text, replacement)
    return text


def _split_page(reader: Any, page_index: int, destination: Path) -> None:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_page(reader.pages[page_index])
    with destination.open("wb") as stream:
        writer.write(stream)


def _make_error_docx(path: Path, page_number: int, error: str) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading(f"OCR error on source page {page_number}", level=1)
    doc.add_paragraph(error)
    doc.save(path)


def _merge_docx(page_files: Iterable[Path], destination: Path) -> None:
    from docx import Document
    from docxcompose.composer import Composer

    files = list(page_files)
    if not files:
        raise RuntimeError("No page-level Word files were produced.")
    master = Document(str(files[0]))
    composer = Composer(master)
    for page_file in files[1:]:
        master.add_page_break()
        composer.append(Document(str(page_file)))
    temp = destination.with_suffix(".tmp.docx")
    composer.save(str(temp))
    os.replace(temp, destination)


def _build_logger(log_path: Path) -> logging.Logger:
    class AppendFileHandler(logging.Handler):
        def emit(self, record: logging.LogRecord) -> None:
            try:
                with log_path.open("a", encoding="utf-8") as stream:
                    stream.write(self.format(record) + "\n")
            except Exception:
                self.handleError(record)

    logger = logging.getLogger(f"scan2doc.{log_path}")
    logger.setLevel(logging.INFO)
    logger.propagate = False
    if not logger.handlers:
        # Opening for each record avoids Windows locking the output directory
        # after a job finishes or fails validation.
        handler = AppendFileHandler()
        handler.setFormatter(
            logging.Formatter("%(asctime)s %(levelname)s %(message)s")
        )
        logger.addHandler(handler)
    return logger


def run_job(
    raw_config: JobConfig,
    progress: ProgressCallback | None = None,
    cancel_event: threading.Event | None = None,
    engine_factory: Callable[[JobConfig], OcrEngine] = PaddleStructureEngine,
) -> JobSummary:
    config = raw_config.normalized()
    validate_input_pdf(config.input_pdf)
    if not config.make_markdown and not config.make_docx:
        raise ValueError("Select at least one output format.")

    config.output_dir.mkdir(parents=True, exist_ok=True)
    work_dir = config.output_dir / ".scan2doc"
    pages_dir = work_dir / "pages"
    page_docx_dir = work_dir / "docx"
    page_md_dir = work_dir / "markdown"
    for folder in (work_dir, pages_dir, page_docx_dir, page_md_dir):
        folder.mkdir(parents=True, exist_ok=True)
    log_path = config.output_dir / "scan2doc.log"
    manifest_path = work_dir / "manifest.json"
    logger = _build_logger(log_path)

    notify = progress or (lambda _message, _done, _total: None)
    notify("Checking the PDF...", 0, 0)
    input_digest = file_sha256(config.input_pdf)
    fingerprint = config_fingerprint(config, input_digest)

    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("pypdf is not installed. Run Install Scan2Doc.bat.") from exc
    reader = PdfReader(str(config.input_pdf), strict=False)
    if reader.is_encrypted:
        try:
            unlocked = reader.decrypt("")
        except Exception as exc:
            raise ValueError("Password-protected PDFs are not supported.") from exc
        if not unlocked:
            raise ValueError("Password-protected PDFs are not supported.")
    total = len(reader.pages)
    if total < 1:
        raise ValueError("The PDF contains no pages.")

    if manifest_path.exists():
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        if manifest.get("fingerprint") != fingerprint:
            raise ValueError(
                "This output folder belongs to a different PDF or settings. "
                "Choose a new output folder to protect existing results."
            )
        if int(manifest.get("total_pages", -1)) != total:
            raise ValueError("The PDF page count changed since the saved job.")
    else:
        manifest = {
            "schema": 1,
            "app_version": APP_VERSION,
            "fingerprint": fingerprint,
            "input_path": str(config.input_pdf),
            "input_sha256": input_digest,
            "total_pages": total,
            "settings": {
                key: str(value) if isinstance(value, Path) else value
                for key, value in asdict(config).items()
            },
            "pages": {},
            "created_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
        }
        _atomic_json(manifest_path, manifest)

    complete_before = sum(
        1
        for value in manifest["pages"].values()
        if value.get("status") == "complete"
    )
    notify(f"Checking saved pages ({complete_before}/{total} already done)...", complete_before, total)
    logger.info("Starting %s; %d pages; fingerprint=%s", config.input_pdf, total, fingerprint)
    engine: OcrEngine | None = None

    failed: list[int] = []
    cancelled = False
    for page_index in range(total):
        page_number = page_index + 1
        key = str(page_number)
        previous = manifest["pages"].get(key, {})
        expected_md = page_md_dir / f"page_{page_number:04d}.md"
        page_word_dir = page_docx_dir / f"page_{page_number:04d}"
        expected_docx = page_word_dir / f"page_{page_number:04d}.docx"
        # Recover a page whose OCR/export finished but the wrapper stopped
        # before committing the manifest (for example, an upstream filename
        # convention changed). This is deliberately limited to the page's
        # private work directory.
        if config.make_docx and not expected_docx.exists() and page_word_dir.is_dir():
            recoverable = list(page_word_dir.glob("*.docx"))
            if len(recoverable) == 1:
                os.replace(recoverable[0], expected_docx)
        artifacts_ok = (
            (not config.make_markdown or expected_md.is_file())
            and (not config.make_docx or expected_docx.is_file())
        )
        if previous.get("status") != "complete" and artifacts_ok:
            manifest["pages"][key] = {
                "status": "complete",
                "attempts": previous.get("attempts", 0),
                "recovered": True,
                "updated_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
            }
            _atomic_json(manifest_path, manifest)
            previous = manifest["pages"][key]
        if previous.get("status") == "complete" and artifacts_ok:
            notify(f"Page {page_number}/{total} already complete", page_number, total)
            continue
        if cancel_event and cancel_event.is_set():
            cancelled = True
            logger.info("Cancellation requested before page %d", page_number)
            break

        page_pdf = pages_dir / f"page_{page_number:04d}.pdf"
        _split_page(reader, page_index, page_pdf)
        if engine is None:
            notify("Loading OCR models...", page_number - 1, total)
            engine = engine_factory(config)
        notify(f"Recognizing page {page_number}/{total}...", page_number - 1, total)
        last_error = ""
        for attempt in range(config.retries + 1):
            try:
                started = time.monotonic()
                result = engine.process_page(page_pdf)
                if config.make_markdown:
                    page_text = _write_markdown_assets(
                        result.markdown, page_number, config.output_dir
                    )
                    expected_md.write_text(page_text.rstrip() + "\n", encoding="utf-8")
                if config.make_docx:
                    page_word_dir.mkdir(parents=True, exist_ok=True)
                    for stale_docx in page_word_dir.glob("*.docx"):
                        stale_docx.unlink(missing_ok=True)
                    # PaddleOCR's save_to_word API takes a directory, despite
                    # older docs describing it as a file-or-directory path.
                    # It may add a page suffix (e.g. page_0001_0.docx), so
                    # normalize that upstream-generated filename afterward.
                    result.save_to_word(save_path=str(page_word_dir))
                    produced_docx = list(page_word_dir.glob("*.docx"))
                    if len(produced_docx) == 1 and produced_docx[0] != expected_docx:
                        os.replace(produced_docx[0], expected_docx)
                    if not expected_docx.is_file():
                        raise RuntimeError(
                            "PaddleOCR did not create exactly one page-level DOCX file."
                        )
                try:
                    result.save_to_json(
                        save_path=str(work_dir / f"page_{page_number:04d}.json")
                    )
                except Exception:
                    logger.warning("Page %d JSON diagnostic export failed", page_number)
                elapsed = round(time.monotonic() - started, 2)
                manifest["pages"][key] = {
                    "status": "complete",
                    "attempts": attempt + 1,
                    "seconds": elapsed,
                    "updated_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
                }
                _atomic_json(manifest_path, manifest)
                logger.info("Page %d complete in %.2fs", page_number, elapsed)
                notify(f"Completed page {page_number}/{total}", page_number, total)
                last_error = ""
                break
            except Exception as exc:
                last_error = f"{type(exc).__name__}: {exc}"
                logger.error(
                    "Page %d attempt %d failed: %s\n%s",
                    page_number,
                    attempt + 1,
                    last_error,
                    traceback.format_exc(),
                )
                if attempt < config.retries:
                    notify(
                        f"Page {page_number} failed; retrying ({attempt + 2}/{config.retries + 1})...",
                        page_number - 1,
                        total,
                    )
        if last_error:
            failed.append(page_number)
            manifest["pages"][key] = {
                "status": "failed",
                "attempts": config.retries + 1,
                "error": last_error,
                "updated_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
            }
            if config.make_markdown:
                expected_md.write_text(
                    f"# OCR error on source page {page_number}\n\n{last_error}\n",
                    encoding="utf-8",
                )
            if config.make_docx:
                page_word_dir.mkdir(parents=True, exist_ok=True)
                _make_error_docx(expected_docx, page_number, last_error)
            _atomic_json(manifest_path, manifest)
        try:
            page_pdf.unlink(missing_ok=True)
        except OSError:
            logger.warning("Could not remove temporary page file %s", page_pdf)

    markdown_path: Path | None = None
    docx_path: Path | None = None
    if not cancelled:
        safe_stem = _safe_name(config.input_pdf.stem, "scanned_document")
        if config.make_markdown:
            markdown_path = config.output_dir / f"{safe_stem}.md"
            sections = []
            for page_number in range(1, total + 1):
                page_file = page_md_dir / f"page_{page_number:04d}.md"
                if page_file.exists():
                    sections.append(
                        f"<!-- source-page: {page_number} -->\n\n"
                        + page_file.read_text(encoding="utf-8").rstrip()
                    )
            temp_md = markdown_path.with_suffix(".tmp.md")
            temp_md.write_text("\n\n---\n\n".join(sections) + "\n", encoding="utf-8")
            os.replace(temp_md, markdown_path)
        if config.make_docx:
            docx_path = config.output_dir / f"{safe_stem}.docx"
            docx_files = [
                page_docx_dir / f"page_{page_number:04d}" / f"page_{page_number:04d}.docx"
                for page_number in range(1, total + 1)
                if (
                    page_docx_dir
                    / f"page_{page_number:04d}"
                    / f"page_{page_number:04d}.docx"
                ).exists()
            ]
            _merge_docx(docx_files, docx_path)

    complete = sum(
        1
        for value in manifest["pages"].values()
        if value.get("status") == "complete"
    )
    all_failed = sorted(
        int(key)
        for key, value in manifest["pages"].items()
        if value.get("status") == "failed"
    )
    manifest["last_run"] = {
        "completed_pages": complete,
        "failed_pages": all_failed,
        "cancelled": cancelled,
        "finished_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
    }
    _atomic_json(manifest_path, manifest)
    logger.info("Finished: complete=%d failed=%s cancelled=%s", complete, all_failed, cancelled)
    notify("Paused safely" if cancelled else "Finished", complete, total)
    return JobSummary(
        total_pages=total,
        completed_pages=complete,
        failed_pages=all_failed,
        skipped_pages=complete_before,
        markdown_path=markdown_path,
        docx_path=docx_path,
        log_path=log_path,
        manifest_path=manifest_path,
        cancelled=cancelled,
    )


def default_output_dir(input_pdf: Path) -> Path:
    return input_pdf.parent / f"{_safe_name(input_pdf.stem, 'scanned_document')}_OCR"


def prepare_offline_models(config: JobConfig) -> None:
    """Initialize the pipeline so its signed-infrastructure model cache is populated."""
    online = JobConfig(**{**asdict(config.normalized()), "offline": False})
    PaddleStructureEngine(online)


def remove_empty_work_dirs(output_dir: Path) -> None:
    """Best-effort cleanup of temporary page PDFs only; never deletes OCR results."""
    pages = output_dir / ".scan2doc" / "pages"
    if pages.is_dir():
        for item in pages.glob("page_*.pdf"):
            item.unlink(missing_ok=True)
        try:
            pages.rmdir()
        except OSError:
            pass
